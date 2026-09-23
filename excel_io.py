"""Importación y exportación en Excel (.xlsx) y CSV.

Dos formatos:
  * matrix: una fila por empleado y una columna por habilidad (cómodo para editar a mano).
  * long:   una fila por evaluación, con fecha (cómodo para integraciones y para conservar fechas).
"""
import csv
import io
from datetime import datetime

from flask import send_file

from db import q, q1
from services import LEVELS, ValidationError, as_int, set_assessment, stale_limit

EMP_COLS = ["Código", "Nombre", "Email", "Departamento", "Puesto"]
LONG_COLS = ["Código", "Empleado", "Email", "Habilidad", "Categoría", "Nivel",
             "Fecha evaluación", "Evaluado por", "Origen", "Desactualizada"]


def _local(dt):
    return dt.astimezone().replace(tzinfo=None) if dt else None


def _data():
    emps = q("""SELECT e.id, e.external_id, e.name, e.email, e.department, COALESCE(r.name,'') AS job_role
                FROM employees e LEFT JOIN job_roles r ON r.id=e.job_role_id ORDER BY lower(e.name)""")
    skills = q("SELECT id, name, category, description FROM skills ORDER BY lower(category), lower(name)")
    cells = {(r["employee_id"], r["skill_id"]): r for r in q("SELECT * FROM assessments")}
    return emps, skills, cells


def _matrix_rows(emps, skills, cells):
    header = EMP_COLS + [s["name"] for s in skills]
    rows = [[e["external_id"] or "", e["name"], e["email"] or "", e["department"], e["job_role"]] +
            [cells[(e["id"], s["id"])]["level"] if (e["id"], s["id"]) in cells else None for s in skills]
            for e in emps]
    return header, rows


def _long_rows(emps, skills, cells, limit):
    by_id = {s["id"]: s for s in skills}
    rows = []
    for e in emps:
        for s in skills:
            a = cells.get((e["id"], s["id"]))
            if a:
                rows.append([e["external_id"] or "", e["name"], e["email"] or "", by_id[s["id"]]["name"],
                             by_id[s["id"]]["category"], a["level"], _local(a["evaluated_at"]),
                             a["evaluated_by"], a["source"], "Sí" if a["evaluated_at"] < limit else "No"])
    return rows


def export(fmt="xlsx", layout="matrix"):
    emps, skills, cells = _data()
    limit = stale_limit()
    stamp = datetime.now().strftime("%Y%m%d")

    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf, delimiter=";")
        if layout == "long":
            w.writerow(LONG_COLS)
            for r in _long_rows(emps, skills, cells, limit):
                r[6] = r[6].isoformat(sep=" ", timespec="minutes") if r[6] else ""
                w.writerow(r)
        else:
            header, rows = _matrix_rows(emps, skills, cells)
            w.writerow(header)
            w.writerows([["" if v is None else v for v in r] for r in rows])
        data = io.BytesIO(("\ufeff" + buf.getvalue()).encode("utf-8"))   # BOM: Excel abre bien los acentos
        name = f"skillmatrix_{'evaluaciones' if layout == 'long' else 'matriz'}_{stamp}.csv"
        return send_file(data, mimetype="text/csv", as_attachment=True, download_name=name)

    if fmt != "xlsx":
        raise ValidationError("Formato no admitido: usa xlsx o csv.")
    from openpyxl import Workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Matriz"
    header, rows = _matrix_rows(emps, skills, cells)
    ws.append(header)
    for r in rows:
        ws.append(r)
    stale_font = Font(italic=True, color="9A5B16")
    for i, e in enumerate(emps, start=2):
        for j, s in enumerate(skills, start=len(EMP_COLS) + 1):
            a = cells.get((e["id"], s["id"]))
            if a and a["evaluated_at"] < limit:
                c = ws.cell(i, j)
                c.font = stale_font
    head_fill, head_font = PatternFill("solid", fgColor="17324D"), Font(bold=True, color="FFFFFF")
    for c in ws[1]:
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 24 if i <= len(EMP_COLS) else 12
    ws.column_dimensions["A"].width = 12
    if skills and rows:
        rng = f"{get_column_letter(len(EMP_COLS) + 1)}2:{get_column_letter(len(header))}{len(rows) + 1}"
        ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=0, start_color="F2F4F7",
                                                          mid_type="num", mid_value=3, mid_color="6FB3C9",
                                                          end_type="num", end_value=5, end_color="17324D"))
    ws.freeze_panes = ws.cell(2, len(EMP_COLS) + 1)

    ws2 = wb.create_sheet("Evaluaciones")
    ws2.append(LONG_COLS)
    for r in _long_rows(emps, skills, cells, limit):
        ws2.append(r)
    for row in ws2.iter_rows(min_row=2, min_col=7, max_col=7):
        for c in row:
            c.number_format = "yyyy-mm-dd hh:mm"
    ws3 = wb.create_sheet("Habilidades")
    ws3.append(["Habilidad", "Categoría", "Descripción"])
    for s in skills:
        ws3.append([s["name"], s["category"], s["description"]])
    ws4 = wb.create_sheet("Escala")
    ws4.append(["Nivel", "Significado"])
    for k, v in LEVELS.items():
        ws4.append([k, v])
    ws4.append([])
    ws4.append(["", "Celda vacía = sin evaluar. Cursiva ámbar en «Matriz» = evaluación desactualizada."])
    for sh, widths in ((ws2, [12, 24, 28, 18, 18, 8, 18, 18, 10, 14]), (ws3, [24, 20, 40]), (ws4, [8, 30])):
        for c in sh[1]:
            c.font = Font(bold=True)
        for i, wdt in enumerate(widths, start=1):
            sh.column_dimensions[get_column_letter(i)].width = wdt
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(out, as_attachment=True, download_name=f"skillmatrix_{stamp}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# -------------------------------------------------------------------- importar
def _read(f):
    name = (f.filename or "").lower()
    categories = {}
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(f, read_only=True, data_only=True)
        ws = wb["Matriz"] if "Matriz" in wb.sheetnames else wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if "Habilidades" in wb.sheetnames:
            for r in list(wb["Habilidades"].iter_rows(values_only=True))[1:]:
                if r and r[0]:
                    categories[str(r[0]).strip().lower()] = str(r[1]).strip() if len(r) > 1 and r[1] else ""
        return rows, categories
    if name.endswith(".csv"):
        text = f.read().decode("utf-8-sig", errors="replace")
        first = text.splitlines()[0] if text else ""
        delim = max(";,\t", key=first.count) if first else ";"
        return list(csv.reader(io.StringIO(text), delimiter=delim)), categories
    raise ValidationError("Formato no admitido. Sube un archivo .xlsx o .csv.")


def _cell(r, idx, key):
    i = idx.get(key)
    v = r[i] if i is not None and i < len(r) else None
    return "" if v is None else str(v).strip()


def _date(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        dt = v
    else:
        try:
            dt = datetime.fromisoformat(str(v).strip().replace("Z", "+00:00"))
        except ValueError:
            raise ValidationError(f"fecha «{v}» no válida (usa AAAA-MM-DD o AAAA-MM-DD HH:MM)")
    if dt.tzinfo is None:
        dt = dt.astimezone()      # sin zona horaria = hora local del servidor
    if dt > datetime.now().astimezone():
        raise ValidationError(f"fecha «{v}» en el futuro")
    return dt


def _employee(r, idx, stats, roles):
    code, name, email = _cell(r, idx, "código"), _cell(r, idx, "nombre") or _cell(r, idx, "empleado"), \
        _cell(r, idx, "email").lower()
    row = (q1("SELECT id FROM employees WHERE external_id=%s", (code,)) if code else None) or \
          (q1("SELECT id FROM employees WHERE lower(email)=%s", (email,)) if email else None) or \
          (q1("SELECT id FROM employees WHERE lower(name)=lower(%s) AND email IS NULL AND external_id IS NULL",
              (name,)) if name and not email and not code else None)
    dept, role = _cell(r, idx, "departamento"), _cell(r, idx, "puesto")
    role_id = None
    if role:
        role_id = roles.get(role.lower())
        if not role_id:
            role_id = q1("INSERT INTO job_roles (name) VALUES (%s) RETURNING id", (role,))["id"]
            roles[role.lower()] = role_id
    if row:
        if "departamento" in idx or "puesto" in idx or name:
            q("""UPDATE employees SET name=COALESCE(NULLIF(%s,''), name), department=COALESCE(%s, department),
                 job_role_id=COALESCE(%s, job_role_id), external_id=COALESCE(external_id, NULLIF(%s,'')),
                 email=COALESCE(email, NULLIF(%s,'')), updated_at=now() WHERE id=%s""",
              (name, dept if "departamento" in idx else None, role_id, code, email, row["id"]))
        return row["id"], False
    if not name:
        raise ValidationError("falta el nombre del empleado")
    new = q1("INSERT INTO employees (external_id, name, email, department, job_role_id) VALUES (%s,%s,%s,%s,%s) RETURNING id",
             (code or None, name, email or None, dept, role_id))
    return new["id"], True


def import_file(f, principal, review_unchanged=False):
    rows, categories = _read(f)
    rows = [r for r in rows if r and any(str(c).strip() for c in r if c is not None)]
    if not rows:
        raise ValidationError("El archivo está vacío.")
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    idx = {h.lower(): i for i, h in enumerate(header) if h}
    is_long = "habilidad" in idx and "nivel" in idx
    if not ({"nombre", "empleado", "email", "código"} & idx.keys()):
        raise ValidationError("No se encuentran columnas para identificar al empleado (Código, Email o Nombre). "
                              "Usa como plantilla un archivo exportado.")
    stats = {"format": "long" if is_long else "matrix", "employees_created": 0, "employees_updated": 0,
             "skills_created": 0, "created": 0, "changed": 0, "reviewed": 0, "unchanged": 0, "skipped": 0,
             "warnings": []}
    roles = {r["name"].lower(): r["id"] for r in q("SELECT id, name FROM job_roles")}
    seen = set()

    def skill_id(name, category=""):
        row = q1("SELECT id FROM skills WHERE lower(name)=lower(%s)", (name,))
        if row:
            return row["id"]
        stats["skills_created"] += 1
        return q1("INSERT INTO skills (name, category) VALUES (%s,%s) RETURNING id",
                  (name, category or categories.get(name.lower()) or "Sin categoría"))["id"]

    fixed = {c.lower() for c in EMP_COLS + LONG_COLS}
    skill_cols = [] if is_long else [(i, h, skill_id(h)) for i, h in enumerate(header) if h and h.lower() not in fixed]

    for n, r in enumerate(rows[1:], start=2):
        try:
            emp_id, created = _employee(r, idx, stats, roles)
        except ValidationError as ex:
            stats["warnings"].append(f"Fila {n}: {ex}.")
            continue
        if emp_id not in seen:
            seen.add(emp_id)
            stats["employees_created" if created else "employees_updated"] += 1
        pairs = ([(_cell(r, idx, "habilidad"), r[idx["nivel"]] if idx["nivel"] < len(r) else None,
                   r[idx["fecha evaluación"]] if "fecha evaluación" in idx and idx["fecha evaluación"] < len(r) else None)]
                 if is_long else [(h, r[i] if i < len(r) else None, None) for i, h, _ in skill_cols])
        for sk, raw, when in pairs:
            if raw is None or str(raw).strip() == "" or not sk:
                continue
            try:
                level = as_int(raw, 0, 5, sk)
                dt = _date(when)
            except ValidationError as ex:
                stats["warnings"].append(f"Fila {n}, «{sk}»: {ex}; se ignora.")
                continue
            sid = skill_id(sk, _cell(r, idx, "categoría"))
            res = set_assessment(emp_id, sid, level, principal.actor, "import", dt,
                                 touch_if_same=review_unchanged or dt is not None)
            stats[res] += 1
    stats["warnings"] = stats["warnings"][:100]
    return stats
