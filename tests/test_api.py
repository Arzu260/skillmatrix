"""Pruebas de la API. Necesitan una base de datos PostgreSQL de pruebas VACÍA (se borra entera):

    set TEST_DATABASE_URL=postgresql://skillmatrix:skillmatrix@localhost:5432/skillmatrix_test
    pytest -q
"""
import io
import os

import psycopg
import pytest

TEST_DB = os.environ.get("TEST_DATABASE_URL")
if not TEST_DB:
    pytest.skip("Define TEST_DATABASE_URL para ejecutar las pruebas.", allow_module_level=True)

with psycopg.connect(TEST_DB, autocommit=True) as c:
    c.execute("DROP SCHEMA public CASCADE; CREATE SCHEMA public;")
os.environ["DATABASE_URL"] = TEST_DB
os.environ["SKILLMATRIX_DEMO"] = "1"
os.environ.pop("SECRET_KEY", None)

from app import app  # noqa: E402

H = {"X-Requested-With": "SkillMatrix"}


@pytest.fixture(scope="module")
def admin():
    c = app.test_client()
    r = c.post("/api/v1/auth/login", json={"username": "admin", "password": "admin"})
    assert r.status_code == 200
    return c


@pytest.fixture(scope="module")
def tokens(admin):
    out = {}
    for scope in ("read", "write"):
        r = admin.post("/api/v1/tokens", json={"name": f"app-{scope}", "scope": scope}, headers=H)
        assert r.status_code == 201
        out[scope] = {"Authorization": f"Bearer {r.json['token']}"}
    return out


def test_requires_auth():
    c = app.test_client()
    assert c.get("/api/v1/employees").status_code == 401
    assert c.get("/api/v1/employees", headers={"Authorization": "Bearer malo"}).status_code == 401
    assert c.post("/api/v1/auth/login", json={"username": "admin", "password": "x"}).status_code == 401


def test_csrf_header_required_for_session_writes(admin):
    r = admin.post("/api/v1/skills", json={"name": "Go"})
    assert r.status_code == 403


def test_read_token_cannot_write(tokens):
    c = app.test_client()
    assert c.get("/api/v1/skills", headers=tokens["read"]).status_code == 200
    assert c.post("/api/v1/skills", json={"name": "Go"}, headers=tokens["read"]).status_code == 403
    assert c.get("/api/v1/users", headers=tokens["write"]).status_code == 403   # tokens nunca son admin


def test_employee_crud_and_filters(tokens):
    c, h = app.test_client(), tokens["write"]
    r = c.post("/api/v1/employees", json={"name": "Laura Vidal", "email": "Laura.Vidal@empresa.com",
                                          "external_id": "E101", "department": "QA", "job_role": "Tester"}, headers=h)
    assert r.status_code == 201, r.json
    emp = r.json
    assert emp["email"] == "laura.vidal@empresa.com" and emp["job_role"] == "Tester"
    assert c.post("/api/v1/employees", json={"name": "X", "email": "laura.vidal@empresa.com"}, headers=h).status_code == 409
    assert c.post("/api/v1/employees", json={"email": "sin@nombre.com"}, headers=h).status_code == 422
    r = c.patch(f"/api/v1/employees/{emp['id']}", json={"department": "Calidad"}, headers=h)
    assert r.json["department"] == "Calidad" and r.json["name"] == "Laura Vidal"
    r = c.get("/api/v1/employees?external_id=E101", headers=h)
    assert r.json["total"] == 1
    r = c.get("/api/v1/employees?limit=3&offset=0", headers=h)
    assert len(r.json["items"]) == 3 and r.json["total"] >= 9
    assert c.delete(f"/api/v1/employees/{emp['id']}", headers=h).status_code == 204
    assert c.get(f"/api/v1/employees/{emp['id']}", headers=h).status_code == 404


def test_bulk_assessments_atomic(tokens):
    c, h = app.test_client(), tokens["write"]
    bad = [{"employee_email": "lucia.martin@empresa.com", "skill": "Kubernetes", "level": 3},
           {"employee_email": "noexiste@empresa.com", "skill": "Java", "level": 2},
           {"employee_external_id": "E002", "skill": "Java", "level": 9}]
    r = c.post("/api/v1/assessments", json={"items": bad}, headers=h)
    assert r.status_code == 422 and len(r.json["errors"]) == 2
    r = c.get("/api/v1/assessments?employee_email=lucia.martin@empresa.com&skill=Kubernetes", headers=h)
    assert r.json["total"] == 0                    # no se guardó nada

    good = [{"employee_email": "lucia.martin@empresa.com", "skill": "Kubernetes", "level": 3},
            {"employee_external_id": "E002", "skill": "java", "level": 3},              # mismo nivel → reviewed
            {"employee_external_id": "E002", "skill": "Terraform", "level": 2, "category": "Infraestructura"}]
    r = c.post("/api/v1/assessments?create_missing_skills=true", json=good, headers=h)
    assert r.status_code == 200, r.json
    assert (r.json["created"], r.json["reviewed"]) == (2, 1)
    r = c.get("/api/v1/assessments?employee_external_id=E002&skill=Java", headers=h)
    item = r.json["items"][0]
    assert item["stale"] is False and item["evaluated_by"] == "api:app-write" and item["source"] == "api"


def test_older_evaluation_is_skipped(tokens):
    c, h = app.test_client(), tokens["write"]
    item = {"employee_external_id": "E003", "skill": "Docker", "level": 1, "evaluated_at": "2020-01-01T00:00:00Z"}
    r = c.post("/api/v1/assessments", json=item, headers=h)
    assert r.json["skipped"] == 1
    r = c.post("/api/v1/assessments", json={**item, "evaluated_at": "2099-01-01"}, headers=h)
    assert r.status_code == 422                     # fecha futura


def test_stale_detection_and_review(tokens):
    c, h = app.test_client(), tokens["write"]
    # Javier Sanz (E004) tiene datos de demo de hace ~20 meses
    emp = c.get("/api/v1/employees?external_id=E004", headers=h).json["items"][0]
    assert emp["stale"] == emp["assessed"] > 0
    assert c.get("/api/v1/assessments?stale=true&employee_external_id=E004", headers=h).json["total"] == emp["assessed"]
    assert any(e["id"] == emp["id"] for e in c.get("/api/v1/employees?has_stale=true", headers=h).json["items"])
    r = c.post(f"/api/v1/employees/{emp['id']}/review", json={}, headers=h)
    assert r.json["reviewed"] == emp["assessed"]
    emp2 = c.get(f"/api/v1/employees/{emp['id']}", headers=h).json
    assert emp2["stale"] == 0 and all(not a["stale"] for a in emp2["assessments"])


def test_single_put_delete_and_history(admin):
    m = admin.get("/api/v1/matrix").json
    emp, sk = m["employees"][0]["id"], m["skills"][0]["id"]
    r = admin.put(f"/api/v1/assessments/{emp}/{sk}", json={"level": 5}, headers=H)
    assert r.status_code == 200 and r.json["level"] == 5 and r.json["evaluated_by"] == "admin"
    assert admin.delete(f"/api/v1/assessments/{emp}/{sk}", headers=H).status_code == 204
    assert f"{emp}-{sk}" not in admin.get("/api/v1/matrix").json["assessments"]
    hist = admin.get(f"/api/v1/employees/{emp}/history").json["items"]
    assert hist[0]["new_level"] is None and hist[1]["new_level"] == 5


def test_roles_and_gaps(tokens):
    c, h = app.test_client(), tokens["write"]
    r = c.post("/api/v1/job-roles", json={"name": "Arquitecto cloud", "requirements": [
        {"skill": "AWS", "required_level": 4}, {"skill": "Kubernetes", "required_level": 3}]}, headers=h)
    assert r.status_code == 201 and len(r.json["requirements"]) == 2
    rid = r.json["id"]
    g = c.get(f"/api/v1/job-roles/{rid}/gaps?scope=all", headers=h).json
    assert len(g["employees"]) >= 8 and {b["skill"] for b in g["by_skill"]} == {"AWS", "Kubernetes"}
    assert g["employees"][0]["coverage"] >= g["employees"][-1]["coverage"]
    assert c.post("/api/v1/job-roles", json={"name": "arquitecto CLOUD"}, headers=h).status_code == 409


def test_settings_change_stale_threshold(admin):
    before = admin.get("/api/v1/dashboard").json["totals"]["stale"]
    assert admin.put("/api/v1/settings", json={"stale_months": 60}, headers=H).json["stale_months"] == 60
    assert admin.get("/api/v1/dashboard").json["totals"]["stale"] < before
    admin.put("/api/v1/settings", json={"stale_months": 12}, headers=H)


def test_export_import_roundtrip(admin):
    from openpyxl import load_workbook
    x = admin.get("/api/v1/export?format=xlsx")
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.data))
    assert {"Matriz", "Evaluaciones", "Habilidades", "Escala"} <= set(wb.sheetnames)
    # reimportar sin cambios no debe renovar fechas
    r = admin.post("/api/v1/import", data={"file": (io.BytesIO(x.data), "m.xlsx")},
                   content_type="multipart/form-data", headers=H)
    assert r.status_code == 200, r.json
    assert r.json["changed"] == 0 and r.json["reviewed"] == 0 and r.json["unchanged"] > 0
    # formato largo con fecha
    csv_data = "Código;Empleado;Email;Habilidad;Nivel;Fecha evaluación\nE007;Ana Pérez;;Go;3;2026-01-15\n"
    r = admin.post("/api/v1/import", data={"file": (io.BytesIO(csv_data.encode()), "e.csv")},
                   content_type="multipart/form-data", headers=H)
    assert r.json["format"] == "long" and r.json["created"] == 1 and r.json["skills_created"] == 1
    a = admin.get("/api/v1/assessments?employee_external_id=E007&skill=Go").json["items"][0]
    assert a["evaluated_at"].startswith("2026-01-1") and a["source"] == "import"
    csv_long = admin.get("/api/v1/export?format=csv&layout=long").data.decode("utf-8-sig")
    assert csv_long.splitlines()[0].startswith("Código;Empleado")


def test_openapi_and_docs():
    c = app.test_client()
    spec = c.get("/api/v1/openapi.json").json
    assert spec["openapi"].startswith("3.") and "/assessments" in spec["paths"]
    assert c.get("/api/docs").status_code == 200
    assert c.get("/api/health").json["status"] == "ok"
